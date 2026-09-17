from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.session import SessionLocal
from app.models.polling_unit import PollingUnit


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parents[2]

EXCEL_FILE = BASE_DIR / "Nasarawa_INEC_Polling_Units.xlsx"

SHEET_NAME = "Polling_Units"

TARGET_PER_PU = 30

NASARAWA_STATE_CODE = "25"

NASARAWA_STATE_NAME = "NASARAWA"


# ============================================================
# HELPERS
# ============================================================

def clean(value):
    """
    Convert Excel values to clean strings.
    """
    if value is None:
        return ""

    return str(value).strip()


# ============================================================
# IMPORT
# ============================================================

def import_polling_units():

    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Excel file not found: {EXCEL_FILE}"
        )

    print("=" * 70)
    print("NASARAWA POLLING UNIT IMPORT")
    print("=" * 70)
    print(f"Source: {EXCEL_FILE}")
    print(f"Target per polling unit: {TARGET_PER_PU}")
    print()

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:

        if SHEET_NAME not in workbook.sheetnames:
            raise RuntimeError(
                f"Worksheet '{SHEET_NAME}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[SHEET_NAME]

        rows = worksheet.iter_rows(values_only=True)

        header = next(rows, None)

        if not header:
            raise RuntimeError(
                "The Excel worksheet is empty."
            )

        headers = [
            clean(value)
            for value in header
        ]

        print("Columns found:")
        print(headers)
        print()

        # ====================================================
        # REQUIRED COLUMNS
        # ====================================================

        required_columns = [
            "state_code",
            "state_name",
            "lga_code",
            "lga_name",
            "ward_code",
            "ward_name",
            "pu_code",
            "pu_name",
            "pu_location",
            "full_code",
            "portal_id",
        ]

        missing = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing:
            raise RuntimeError(
                f"Missing required columns: {missing}"
            )

        index = {
            column: headers.index(column)
            for column in required_columns
        }

        # ====================================================
        # DATABASE
        # ====================================================

        db: Session = SessionLocal()

        inserted = 0
        updated = 0
        skipped = 0

        seen_full_codes = set()

        try:

            # =================================================
            # PROCESS EXCEL
            # =================================================

            for row_number, row in enumerate(
                rows,
                start=2,
            ):

                if not row:
                    skipped += 1
                    continue

                # ---------------------------------------------
                # STATE
                # ---------------------------------------------

                state_code = clean(
                    row[index["state_code"]]
                )

                state_name = clean(
                    row[index["state_name"]]
                )

                # Only import Nasarawa
                if (
                    state_code != NASARAWA_STATE_CODE
                    and
                    state_name.upper()
                    != NASARAWA_STATE_NAME
                ):
                    continue

                # ---------------------------------------------
                # LOCATION
                # ---------------------------------------------

                lga_code = clean(
                    row[index["lga_code"]]
                )

                lga_name = clean(
                    row[index["lga_name"]]
                )

                ward_code = clean(
                    row[index["ward_code"]]
                )

                ward_name = clean(
                    row[index["ward_name"]]
                )

                pu_code = clean(
                    row[index["pu_code"]]
                )

                pu_name = clean(
                    row[index["pu_name"]]
                )

                pu_location = clean(
                    row[index["pu_location"]]
                )

                full_code = clean(
                    row[index["full_code"]]
                )

                # ---------------------------------------------
                # VALIDATE FULL CODE
                # ---------------------------------------------

                if not full_code:

                    skipped += 1

                    print(
                        f"Skipping row {row_number}: "
                        f"missing full_code"
                    )

                    continue

                # ---------------------------------------------
                # DUPLICATE EXCEL ROW
                # ---------------------------------------------

                if full_code in seen_full_codes:

                    skipped += 1

                    print(
                        f"Skipping row {row_number}: "
                        f"duplicate full_code {full_code}"
                    )

                    continue

                seen_full_codes.add(full_code)

                # ---------------------------------------------
                # PORTAL ID
                # ---------------------------------------------

                portal_id_value = row[
                    index["portal_id"]
                ]

                portal_id = None

                if portal_id_value not in (
                    None,
                    "",
                ):

                    try:

                        portal_id = int(
                            float(
                                portal_id_value
                            )
                        )

                    except (
                        ValueError,
                        TypeError,
                    ):

                        portal_id = None

                # ---------------------------------------------
                # FIND EXISTING POLLING UNIT
                # ---------------------------------------------

                polling_unit = (
                    db.query(PollingUnit)
                    .filter(
                        PollingUnit.full_code
                        == full_code
                    )
                    .first()
                )

                # =================================================
                # UPDATE EXISTING
                # =================================================

                if polling_unit:

                    # -----------------------------------------
                    # IMPORTANT:
                    #
                    # Do NOT reset registered_count.
                    # -----------------------------------------

                    polling_unit.state_code = (
                        state_code
                    )

                    polling_unit.state_name = (
                        state_name
                    )

                    polling_unit.lga_code = (
                        lga_code
                    )

                    polling_unit.lga_name = (
                        lga_name
                    )

                    polling_unit.ward_code = (
                        ward_code
                    )

                    polling_unit.ward_name = (
                        ward_name
                    )

                    polling_unit.pu_code = (
                        pu_code
                    )

                    polling_unit.pu_name = (
                        pu_name
                    )

                    polling_unit.pu_location = (
                        pu_location
                    )

                    polling_unit.portal_id = (
                        portal_id
                    )

                    # -----------------------------------------
                    # Preserve existing target.
                    # If NULL/0, restore default 30.
                    # -----------------------------------------

                    if not polling_unit.registration_target:

                        polling_unit.registration_target = (
                            TARGET_PER_PU
                        )

                    # -----------------------------------------
                    # Preserve registered count
                    # -----------------------------------------

                    registered = (
                        polling_unit.registered_count
                        or 0
                    )

                    polling_unit.registered_count = (
                        registered
                    )

                    # -----------------------------------------
                    # Recalculate status
                    # -----------------------------------------

                    polling_unit.status = (
                        "FULL"
                        if registered
                        >= polling_unit.registration_target
                        else "OPEN"
                    )

                    updated += 1

                # =================================================
                # INSERT NEW
                # =================================================

                else:

                    polling_unit = PollingUnit(

                        state_code=state_code,

                        state_name=state_name,

                        lga_code=lga_code,

                        lga_name=lga_name,

                        ward_code=ward_code,

                        ward_name=ward_name,

                        pu_code=pu_code,

                        pu_name=pu_name,

                        pu_location=pu_location,

                        full_code=full_code,

                        portal_id=portal_id,

                        registration_target=(
                            TARGET_PER_PU
                        ),

                        registered_count=0,

                        status="OPEN",
                    )

                    db.add(polling_unit)

                    inserted += 1

                # =================================================
                # BATCH COMMIT
                # =================================================

                processed = (
                    inserted + updated
                )

                if (
                    processed > 0
                    and processed % 500 == 0
                ):

                    db.commit()

                    print(
                        f"Processed: {processed:,} | "
                        f"Inserted: {inserted:,} | "
                        f"Updated: {updated:,}"
                    )

            # =====================================================
            # FINAL COMMIT
            # =====================================================

            db.commit()

            # =====================================================
            # FINAL STATISTICS
            # =====================================================

            print()
            print("=" * 70)
            print("IMPORT COMPLETE")
            print("=" * 70)

            # -----------------------------------------------------
            # Total Nasarawa polling units
            # -----------------------------------------------------

            nasarawa_units = (
                db.query(PollingUnit)
                .filter(
                    PollingUnit.state_code
                    == NASARAWA_STATE_CODE
                )
                .all()
            )

            nasarawa_count = len(
                nasarawa_units
            )

            # -----------------------------------------------------
            # LGAs
            # -----------------------------------------------------

            lga_rows = (
                db.query(
                    PollingUnit.lga_code
                )
                .filter(
                    PollingUnit.state_code
                    == NASARAWA_STATE_CODE
                )
                .distinct()
                .all()
            )

            lga_count = len(lga_rows)

            # -----------------------------------------------------
            # WARDS / REGISTRATION AREAS
            #
            # IMPORTANT:
            #
            # ward_code alone is NOT enough because codes repeat
            # across LGAs.
            #
            # We therefore use:
            #
            # (lga_code, ward_code)
            # -----------------------------------------------------

            ward_rows = (
                db.query(
                    PollingUnit.lga_code,
                    PollingUnit.ward_code,
                )
                .filter(
                    PollingUnit.state_code
                    == NASARAWA_STATE_CODE
                )
                .distinct()
                .all()
            )

            ward_count = len(
                ward_rows
            )

            # -----------------------------------------------------
            # OPEN POLLING UNITS
            # -----------------------------------------------------

            open_count = (
                db.query(PollingUnit)
                .filter(
                    PollingUnit.state_code
                    == NASARAWA_STATE_CODE,

                    PollingUnit.status
                    == "OPEN",
                )
                .count()
            )

            # -----------------------------------------------------
            # FULL POLLING UNITS
            # -----------------------------------------------------

            full_count = (
                db.query(PollingUnit)
                .filter(
                    PollingUnit.state_code
                    == NASARAWA_STATE_CODE,

                    PollingUnit.status
                    == "FULL",
                )
                .count()
            )

            # -----------------------------------------------------
            # TOTAL REGISTERED
            # -----------------------------------------------------

            total_registered = sum(
                (
                    pu.registered_count
                    or 0
                )
                for pu in nasarawa_units
            )

            # -----------------------------------------------------
            # TOTAL CAPACITY
            #
            # Use actual target stored on each PU.
            # -----------------------------------------------------

            total_capacity = sum(
                (
                    pu.registration_target
                    or TARGET_PER_PU
                )
                for pu in nasarawa_units
            )

            # -----------------------------------------------------
            # REMAINING CAPACITY
            # -----------------------------------------------------

            remaining_capacity = max(
                total_capacity
                - total_registered,
                0,
            )

            # =====================================================
            # OUTPUT
            # =====================================================

            print(
                f"Inserted:             {inserted:,}"
            )

            print(
                f"Updated:              {updated:,}"
            )

            print(
                f"Skipped:              {skipped:,}"
            )

            print()

            print(
                f"Nasarawa PUs:         {nasarawa_count:,}"
            )

            print(
                f"LGAs:                 {lga_count:,}"
            )

            print(
                f"Wards:                {ward_count:,}"
            )

            print(
                f"Open PUs:             {open_count:,}"
            )

            print(
                f"Full PUs:             {full_count:,}"
            )

            print(
                f"Registered:           {total_registered:,}"
            )

            print(
                f"Capacity:             {total_capacity:,}"
            )

            print(
                f"Remaining:            {remaining_capacity:,}"
            )

            print()

            # =====================================================
            # VALIDATION
            # =====================================================

            if nasarawa_count != 3256:

                print(
                    f"WARNING: Expected 3,256 PUs, "
                    f"but found {nasarawa_count:,}."
                )

            else:

                print(
                    "✓ Polling unit count verified: 3,256"
                )

            if lga_count != 13:

                print(
                    f"WARNING: Expected 13 LGAs, "
                    f"but found {lga_count:,}."
                )

            else:

                print(
                    "✓ LGA count verified: 13"
                )

            if ward_count != 147:

                print(
                    f"WARNING: Expected 147 wards, "
                    f"but found {ward_count:,}."
                )

            else:

                print(
                    "✓ Ward/Registration Area count verified: 147"
                )

            expected_capacity = (
                3256 * TARGET_PER_PU
            )

            if (
                total_capacity
                != expected_capacity
            ):

                print(
                    f"WARNING: Expected capacity "
                    f"{expected_capacity:,}, "
                    f"but found "
                    f"{total_capacity:,}."
                )

            else:

                print(
                    "✓ Capacity verified: 97,680"
                )

            print()
            print("=" * 70)

        except Exception:

            db.rollback()

            raise

        finally:

            db.close()

    finally:

        workbook.close()


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    import_polling_units()